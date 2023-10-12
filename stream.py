import streamlit as st
import openpyxl
from docx import Document
import os
import shutil
import tempfile
import zipfile


# Streamlit Configuration
st.set_page_config(
    page_title="MSAmagi",
    page_icon=":globe_with_meridians:",
    layout="centered",
    initial_sidebar_state="expanded",
)

st.markdown("<h1 style='color: teal; font-family: Arial, sans-serif; font-size: 24px; font-weight: bold; text-align: center;'>MSA-hantering</h1>", unsafe_allow_html=True)

# Initialize MAIN_TEMP_DIR for the session
if 'MAIN_TEMP_DIR' not in st.session_state:
    st.session_state.MAIN_TEMP_DIR = tempfile.mkdtemp()
    
MSA_DIRECTORY = None
RESULTS_DIRECTORY = None
MAIN_TEMP_DIR = st.session_state.MAIN_TEMP_DIR

# Upload File
uploaded_file = st.file_uploader("Upload Files", type=['zip'])

if uploaded_file:
    MSA_DIRECTORY = os.path.join(MAIN_TEMP_DIR, "uploaded_files")
    os.makedirs(MSA_DIRECTORY, exist_ok=True)
    with zipfile.ZipFile(uploaded_file, 'r') as zip_ref:
        zip_ref.extractall(MSA_DIRECTORY)

    RESULTS_DIRECTORY = os.path.join(MAIN_TEMP_DIR, "results")
    os.makedirs(RESULTS_DIRECTORY, exist_ok=True)

    RESULTS_DIRECTORY = os.path.join(MAIN_TEMP_DIR, "results")
    os.makedirs(RESULTS_DIRECTORY, exist_ok=True)

saved_template_paths = {}

def handle_uploaded_template(uploaded_file, key):
    """Save the uploaded template directly to a file."""
    extension = ".docx" if "Avtalsmall" in key else ".xlsx"
    template_dir = os.path.join(st.session_state.MAIN_TEMP_DIR, "templates")
    os.makedirs(template_dir, exist_ok=True)
    temp_file_path = os.path.join(template_dir, f"{key}{extension}")
    
    with open(temp_file_path, "wb") as f:
        f.write(uploaded_file.read())

    
    # Save only the filename to the session state
    if 'uploaded_templates' not in st.session_state:
        st.session_state.uploaded_templates = {}

    st.session_state.uploaded_templates[key] = f"{key}{extension}"

st.markdown("<h2 style='color: teal; font-family: Arial, sans-serif; font-size: 24px; font-weight: bold; text-align: left;'>Ladda upp mallar nedan</h2>", unsafe_allow_html=True)

# Use containers to visually separate the clients
for customer in ["3GIS", "Hi3G", "Telia"]:
    with st.container():
        st.write(f"**{customer}**")
        for template_type in ["Avtalsmall", "Prisberakningsmall"]:
            key = f"{customer}_{template_type}"

            # Only display file uploader if the template hasn't been uploaded yet
            if key not in st.session_state.get('uploaded_templates', {}):
                uploaded_template = st.file_uploader(f"Ladda upp {template_type} för {customer}", key=key)
                if uploaded_template:
                    handle_uploaded_template(uploaded_template, key)


def get_template_path(customer, template_type):
    """Retrieve the saved template path."""
    key = f"{customer}_{template_type}"
    # Generate path on-the-fly
    if key in st.session_state.uploaded_templates:
        return os.path.join(MAIN_TEMP_DIR, "templates", st.session_state.uploaded_templates[key])
    else:
        st.write(f"Please upload the {template_type} for {customer}.")
        return None


def get_tenant_description(sheet, customer_name, cell_d4_value):
    customer_name = customer_name.lower() if isinstance(customer_name, str) else str(customer_name)
    cell_d4_value = cell_d4_value.lower() if isinstance(cell_d4_value, str) else str(cell_d4_value)
    if 'telia' in customer_name or 'telia sverige ab' in cell_d4_value:
        return sheet["D6"].value
    return sheet["D8"].value

def get_landlord_description(sheet, identified_customer):  # added identified_customer parameter
    if identified_customer == 'Telia':
        return sheet["E7"].value or ""
    else:
        return sheet["E9"].value or ""

def get_coordinate(coord, alternative_coord):
    return coord or alternative_coord

def read_msa_info(msa_path):
    wb = openpyxl.load_workbook(msa_path, data_only=True)
    sheet = wb.active
    identified_customer = identify_customer(sheet["D3"].value, sheet["D4"].value)  # Identifying the customer

    data = {
        'customer_name': sheet["D3"].value,
        'tenant_desc': get_tenant_description(sheet, sheet["D3"].value, sheet["D4"].value),
        'landlord_desc': get_landlord_description(sheet, identified_customer),  # Fetch landlord description based on identified customer
        'x_coord': get_coordinate(sheet["D10"].value, sheet["D9"].value),
        'y_coord': get_coordinate(sheet["D11"].value, sheet["D9"].value),
        'mast_info': sheet["E12"].value,
        'cell_d4_value': sheet["D4"].value
    }
    wb.close()
    return data


def identify_customer(customer_name, cell_d4_value):
    customer_name = customer_name.lower() if isinstance(customer_name, str) else str(customer_name)
    cell_d4_value = cell_d4_value.lower() if isinstance(cell_d4_value, str) else str(cell_d4_value)
    if '3gis' in customer_name:
        return '3GIS'
    elif 'hi3g' in customer_name or 'hi3g access ab' in customer_name:
        return 'Hi3G'
    elif 'telia' in customer_name or 'telia sverige ab' in cell_d4_value:
        return 'Telia'
    return None

def replace_placeholders_in_document(doc, data):
    placeholders = {
        "--::HYRESVARD::--": data['landlord_desc'] or "",
        "--::HYRESGAST::--": data['tenant_desc'] or "",
        "--::x_coordinate::--": str(data['x_coord']) if data['x_coord'] else "",
        "--::y_coordinate::--": str(data['y_coord']) if data['y_coord'] else "",
        "--::masttyp_hojd::--": str(data['mast_info']) if data['mast_info'] else "",
    }
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            for key, value in placeholders.items():
                if key in run.text:
                    run.text = run.text.replace(key, value)

def sanitize_file_name(name):
    sanitized_name = '_'.join(name.split())
    sanitized_name = sanitized_name.replace("/", "_")
    return sanitized_name

def update_sheet_based_on_customer(sheet, customer, landlord_desc, tenant_desc):
    if customer == 'Hi3G':
        sheet['G1'] = landlord_desc
    elif customer == '3GIS':
        sheet['D6'] = landlord_desc
    elif customer == 'Telia':
        if tenant_desc:
            sheet['D7'] = tenant_desc
        sheet['D9'] = landlord_desc

def update_word_template(customer, msa_data):
    #st.write(f"Checking for Word template of {customer} at: {get_template_path(customer, 'Avtalsmall')}")
    template_path = get_template_path(customer, "Avtalsmall")
    if template_path and os.path.exists(template_path):
        doc = Document(template_path)
        replace_placeholders_in_document(doc, msa_data)
        result_folder = os.path.join(RESULTS_DIRECTORY, customer)
        os.makedirs(result_folder, exist_ok=True)
        sanitized_landlord_desc = sanitize_file_name(msa_data['landlord_desc'])
        output_filename = f"{sanitized_landlord_desc}_Avtalsmall.docx"
        output_filepath = os.path.join(result_folder, output_filename)
        doc.save(output_filepath)
        #st.write(f"Generated document for {customer} in {result_folder}")
    else:
        st.write(f"Template for {customer} not found or not uploaded. Check the path: {template_path}")

def update_prisberakningsmall(customer, landlord_desc, tenant_desc, data):
    #st.write(f"Checking for Prisberakningsmall of {customer} at: {get_template_path(customer, 'Prisberakningsmall')}")
    template_path = get_template_path(customer, "Prisberakningsmall")
    try:
        if template_path and os.path.exists(template_path):
            #st.write(f"Retrieved path for {customer}'s Prisberäkningsmall: {template_path}")

            wb = openpyxl.load_workbook(template_path)
            sheet = wb.active

            update_sheet_based_on_customer(sheet, customer, landlord_desc, tenant_desc)  # <-- This line is added.

            result_folder = os.path.join(RESULTS_DIRECTORY, customer)
            os.makedirs(result_folder, exist_ok=True)

            result_file_name = f"{sanitize_file_name(landlord_desc)}_{sanitize_file_name(tenant_desc)}_{customer}_Prisberakningsmall.xlsx"
            result_path = os.path.join(result_folder, result_file_name)

            wb.save(result_path)
            wb.close()
        else:
            st.write(f"The file for {customer}'s Prisberäkningsmall does NOT exist at {template_path}.")
    except Exception as e:
        st.write(f"Error while processing Prisberäkningsmall for {customer}: {str(e)}")


def process_msa_files(msa_directory):
    files_processed = False
    for filename in os.listdir(msa_directory):
        if filename.endswith(".xlsx"):
            msa_path = os.path.join(msa_directory, filename)
            #st.write(f"Processing MSA file: {filename}")
            msa_data = read_msa_info(msa_path)

             # Log the extracted data
            #st.write(f"Data from {filename}: {msa_data}")
            
            customer = identify_customer(msa_data['customer_name'], msa_data['cell_d4_value'])
            if customer:
                update_word_template(customer, msa_data)
                update_prisberakningsmall(customer, msa_data['landlord_desc'], msa_data['tenant_desc'], msa_data)
                files_processed = True
    return files_processed

def compress_results_to_zip(results_directory, zip_filename):
    """Compress all the results into a single .zip file."""
    with zipfile.ZipFile(zip_filename, 'w') as zipf:
        for foldername, subfolders, filenames in os.walk(results_directory):
            for filename in filenames:
                file_path = os.path.join(foldername, filename)
                arcname = os.path.relpath(file_path, results_directory)
                zipf.write(file_path, arcname)

if st.button("Hantera filer"):
    with st.spinner('Hanterar filer...'):
        # Check if all templates are uploaded before processing
        all_templates_provided = True
        for customer in ["3GIS", "Hi3G", "Telia"]:
            for template_type in ["Avtalsmall", "Prisberakningsmall"]:
                key = f"{customer}_{template_type}"
                if key not in st.session_state.get('uploaded_templates', {}).keys():
                    st.write(f"Please upload the {template_type} for {customer}.")
                    all_templates_provided = False
        
        if all_templates_provided:
            files_processed = process_msa_files(MSA_DIRECTORY)

            
        if files_processed:
            zip_filename = os.path.join(MAIN_TEMP_DIR, "processed_results.zip")
            compress_results_to_zip(RESULTS_DIRECTORY, zip_filename)
  ###############   

            with open(zip_filename, "rb") as f:
                zip_bytes = f.read()
                    
            st.download_button(label="Ladda hem hanterade dokument", data=zip_bytes, mime="application/zip")
    